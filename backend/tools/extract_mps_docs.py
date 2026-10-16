r"""从 MPS 只读资料目录提取 docx/xlsx 文本，供编制 check-list 使用（只读，不修改源目录）。

源目录默认 F:\MPS，可用环境变量 MPS_DIR 覆盖。
用法：在 backend 目录执行  python tools/extract_mps_docs.py
输出：backend/.tmp/mps-extract/<原名>.md
"""
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

MPS_DIR = Path(os.environ.get("MPS_DIR", r"F:\MPS"))
OUT_DIR = Path(__file__).resolve().parents[1] / ".tmp" / "mps-extract"


def extract_docx(path: Path, out: Path) -> None:
    import docx

    doc = docx.Document(str(path))
    lines: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for ti, table in enumerate(doc.tables):
        lines.append(f"\n[表格 {ti + 1}]")
        for row in table.rows:
            lines.append(" | ".join(c.text.replace("\n", " ") for c in row.cells))
    out.write_text("\n".join(lines), encoding="utf-8")


def extract_xlsx(path: Path, out: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"\n===== 工作表: {ws.title} (rows={ws.max_row}, cols={ws.max_column}) =====")
        for row in ws.iter_rows(values_only=True):
            if all(v is None for v in row):
                continue
            lines.append(" | ".join("" if v is None else str(v) for v in row))
    out.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for f in MPS_DIR.glob("*.docx"):
        extract_docx(f, OUT_DIR / (f.stem + ".md"))
        print("extracted:", f.name)
    for f in MPS_DIR.glob("*.xlsx"):
        extract_xlsx(f, OUT_DIR / (f.stem + ".md"))
        print("extracted:", f.name)


if __name__ == "__main__":
    main()
